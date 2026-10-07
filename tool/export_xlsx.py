"""
Noodexport: de draft tot nu toe als Excel, in de vorm van de oude draftsheet.

    py tool\\export_xlsx.py                     -> exports\\LoW draft <datum> <uur>.xlsx
    py tool\\export_xlsx.py --open              en meteen openen (zo draait export.bat)
    py tool\\export_xlsx.py --template "2025 Fantasy Auction Draft.xlsx"

Waarvoor: als de app vastloopt of onhandig wordt tijdens de draft, moet de stand
eruit zonder de app. Daarom leest dit script data\\draft.db rechtstreeks, read-only.
Elke pick staat daar vast zodra hij gemaakt is, dus het werkt ook als de server
hangt, gecrasht is of nooit gestart werd. Het raakt de app en de database niet.

Layout (de variant van 2019 tot 2025, zoals tool/parse_history.py die leest):
per team twee kolommen, prijs links en speler rechts. De teamnaam staat in de
spelerskolom, de rij eronder heeft 200 in de prijskolom, daaronder de picks.
Onderaan per team Uitgegeven, Resterend en Max bid als formules, zodat je in
Excel verder kan draften als het moet. Een tweede tabblad "Picks" geeft de picks
in volgorde.

Met --template vult het script een kopie van een bestaande draftsheet: het
neemt het laatste tabblad "Auction <jaar>" als voorbeeld, kopieert het naar een
nieuw tabblad voor dit jaar, zet de huidige teamnamen in de kolommen en
vervangt de picks. Opmaak en formules van het voorbeeld blijven staan.
Zonder --template zoekt het data\\draft_template.xlsx; is die er niet, dan
bouwt het de sheet zelf op.

Vereist openpyxl (staat in requirements.txt).
"""

import argparse
import json
import os
import re
import sqlite3
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("openpyxl ontbreekt. Installeer met: py -m pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_history import LABEL, YEAR, budget_row  # noqa: E402

DEFAULT_DB = os.path.join(ROOT, "data", "draft.db")
DEFAULT_TEMPLATE = os.path.join(ROOT, "data", "draft_template.xlsx")
DEFAULT_OUT = os.path.join(ROOT, "exports")
CONFIG = os.path.join(ROOT, "tool", "config.json")


# ---------------------------------------------------------------- lezen

def read_draft(db_path):
    """Teams, picks en blok uit de database, zonder er iets in te schrijven."""
    if not os.path.exists(db_path):
        raise SystemExit(f"Database niet gevonden: {db_path}")
    uri = "file:" + os.path.abspath(db_path).replace("\\", "/") + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row
    try:
        settings = {r["k"]: json.loads(r["v"]) for r in con.execute("SELECT k, v FROM settings")}
        picks = [dict(r) for r in con.execute("SELECT * FROM picks ORDER BY seq")]
    finally:
        con.close()
    teams = list(settings.get("teams") or [])
    for p in picks:                 # een pick van een team dat niet meer in de lijst staat
        if p["team"] not in teams:
            teams.append(p["team"])
    return {"teams": teams, "picks": picks, "block": settings.get("block"),
            "me": settings.get("me")}


def read_league(config_path=CONFIG):
    league = {"budget": 200, "roster_spots": 13}
    try:
        with open(config_path, encoding="utf-8") as f:
            cfg = json.load(f)
        league.update({k: int(cfg[k]) for k in league if k in cfg})
    except (OSError, ValueError):
        pass
    return league


def by_team(draft):
    out = {t: [] for t in draft["teams"]}
    for p in draft["picks"]:
        out[p["team"]].append(p)
    return out


# ---------------------------------------------------------------- zelf opbouwen

THIN = Side(style="thin", color="999999")
HEAD = PatternFill("solid", fgColor="1F3864")
BUDGET = PatternFill("solid", fgColor="D9E1F2")
FOOT = PatternFill("solid", fgColor="F2F2F2")


def build_sheet(ws, draft, league, now):
    """De draftsheet zelf opbouwen: titel, teamrij, budgetrij, picks en totalen."""
    spots, budget = league["roster_spots"], league["budget"]
    rosters = by_team(draft)
    hdr, brow = 2, 3
    first, last = brow + 1, brow + spots
    # Een team kan in de app niet meer dan `spots` picks hebben, maar een
    # database die met andere instellingen gevuld werd wel: dan schuift de voet op.
    last = max(last, brow + max((len(v) for v in rosters.values()), default=0))
    foot = last + 1

    ws["A1"] = f"LoW auction draft {now:%Y} (export {now:%d/%m/%Y %H:%M}, {len(draft['picks'])} picks)"
    ws["A1"].font = Font(bold=True, size=12)
    for i, team in enumerate(draft["teams"]):
        pc, nc = 2 * i + 1, 2 * i + 2
        P, N = get_column_letter(pc), get_column_letter(nc)
        ws.column_dimensions[P].width = 6
        ws.column_dimensions[N].width = 24

        ws.cell(hdr, pc).fill = HEAD
        c = ws.cell(hdr, nc, team)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        ws.cell(brow, pc, budget).fill = BUDGET
        ws.cell(brow, nc, "Budget").fill = BUDGET

        for r, p in enumerate(rosters[team], start=first):
            ws.cell(r, pc, int(p["price"]))
            ws.cell(r, nc, p["name"])
        for r in range(first, last + 1):
            ws.cell(r, pc).border = Border(left=THIN, bottom=THIN)
            ws.cell(r, nc).border = Border(right=THIN, bottom=THIN)

        picks = f"{P}{first}:{P}{last}"
        rest, open_ = f"{P}{foot + 1}", f"({spots}-COUNT({picks}))"
        rows = [("Uitgegeven", f"=SUM({picks})"),
                ("Resterend", f"={P}{brow}-{P}{foot}"),
                ("Max bid", f"=IF({open_}<=0,0,MAX(0,{rest}-({open_}-1)))")]
        for k, (label, formula) in enumerate(rows):
            ws.cell(foot + k, pc, formula).fill = FOOT
            lc = ws.cell(foot + k, nc, label)
            lc.fill, lc.font = FOOT, Font(italic=True)

    ws.freeze_panes = ws.cell(first, 1)
    if draft["block"]:
        b = draft["block"]
        ws.cell(foot + 4, 1, f"Op het blok (niet verkocht): {b.get('name') or b.get('player_id')}"
                             f", genomineerd door {b.get('nominator')}").font = Font(italic=True)


# ---------------------------------------------------------------- template vullen

def template_columns(ws, brow):
    """(prijskolom, naamkolom) per team in een bestaande sheet, links naar rechts.

    Dezelfde twee oriëntaties als parse_history: teamnaam in de naamkolom met de
    prijs links ervan, of teamnaam in de prijskolom met de naam rechts ervan."""
    best = []
    for shift in (0, 1):
        cols = []
        for c in range(1, ws.max_column + 1):
            nm = ws.cell(brow - 1, c).value
            if not (nm and str(nm).strip()):
                continue
            pc, nc = (c - 1, c) if shift == 0 else (c, c + 1)
            if pc >= 1 and nc <= ws.max_column and ws.cell(brow, pc).value in (200, 200.0):
                cols.append((pc, nc))
        if len(cols) > len(best):
            best = cols
    return best


def pick_rows(ws, brow, nc, spots):
    """Rijen voor de picks: vanaf onder de budgetrij tot de eerste labelrij."""
    r = brow + 1
    while r <= ws.max_row:
        v = ws.cell(r, nc).value
        if v is not None and LABEL.match(str(v).strip()):
            return list(range(brow + 1, r))
        r += 1
    return list(range(brow + 1, brow + 1 + spots))


def fill_template(path, draft, league, now):
    wb = openpyxl.load_workbook(path)
    auctions = [ws for ws in wb.worksheets
                if ws.title.startswith("Auction") and YEAR.search(ws.title) and budget_row(ws)]
    if not auctions:
        raise SystemExit(f"Geen tabblad 'Auction <jaar>' met een budgetrij in {path}")
    src = max(auctions, key=lambda ws: YEAR.search(ws.title).group(1))
    title = f"Auction {now:%Y}"
    if src.title != title:
        if title in wb.sheetnames:
            del wb[title]
        ws = wb.copy_worksheet(src)
        ws.title = title
        wb.move_sheet(ws, -(len(wb.sheetnames) - 1))     # vooraan
        wb.active = 0
    else:
        ws = src

    brow = budget_row(ws)
    cols = template_columns(ws, brow)
    if len(cols) < len(draft["teams"]):
        raise SystemExit(f"Het voorbeeld heeft {len(cols)} teamkolommen, de draft "
                         f"{len(draft['teams'])} teams. Gebruik de export zonder --template.")
    rosters = by_team(draft)
    for i, (pc, nc) in enumerate(cols):
        team = draft["teams"][i] if i < len(draft["teams"]) else None
        hc = nc if ws.cell(brow - 1, nc).value else pc
        ws.cell(brow - 1, hc).value = team
        rows = pick_rows(ws, brow, nc, league["roster_spots"])
        picks = rosters.get(team, [])
        if len(picks) > len(rows):
            raise SystemExit(f"{team} heeft {len(picks)} picks, het voorbeeld maar {len(rows)} rijen.")
        for r in rows:
            ws.cell(r, pc).value = ws.cell(r, nc).value = None
        for r, p in zip(rows, picks):
            ws.cell(r, pc).value = int(p["price"])
            ws.cell(r, nc).value = p["name"]
    return wb, ws


# ---------------------------------------------------------------- picks-tabblad

def picks_sheet(ws, draft):
    ws.append(["#", "Team", "Speler", "Prijs", "Bron", "Tijd"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, p in enumerate(draft["picks"], start=1):
        ws.append([i, p["team"], p["name"], int(p["price"]), p.get("source"),
                   datetime.fromtimestamp(p["ts"]).strftime("%H:%M:%S")])
    for col, w in zip("ABCDEF", (5, 20, 26, 7, 8, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    for c in ws["D"][1:]:
        c.alignment = Alignment(horizontal="right")


# ---------------------------------------------------------------- main

def export(db_path, out_dir, template=None, now=None, config_path=CONFIG):
    now = now or datetime.now()
    draft = read_draft(db_path)
    league = read_league(config_path)
    if template:
        wb, _ = fill_template(template, draft, league, now)
        name = "Picks " + now.strftime("%Y")
        if name in wb.sheetnames:
            del wb[name]
        picks_sheet(wb.create_sheet(name, 1), draft)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Auction {now:%Y}"
        build_sheet(ws, draft, league, now)
        picks_sheet(wb.create_sheet("Picks"), draft)

    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"LoW draft {now:%Y-%m-%d %H%M%S}.xlsx")
    wb.save(path)
    return path, draft


def main(argv=None):
    ap = argparse.ArgumentParser(description="Exporteer de draft naar Excel, zonder de app.")
    ap.add_argument("--db", default=os.environ.get("LOW_DB", DEFAULT_DB))
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--template", default=None,
                    help="bestaande draftsheet om te vullen (default data\\draft_template.xlsx als die bestaat)")
    ap.add_argument("--open", action="store_true", help="het bestand meteen openen")
    a = ap.parse_args(argv)

    template = a.template or (DEFAULT_TEMPLATE if os.path.exists(DEFAULT_TEMPLATE) else None)
    if template and not os.path.exists(template):
        raise SystemExit(f"Template niet gevonden: {template}")
    path, draft = export(a.db, a.out, template)
    print(f"{len(draft['picks'])} picks over {len(draft['teams'])} teams"
          f"{' (template: ' + os.path.basename(template) + ')' if template else ''}")
    print(f"geschreven naar {path}")
    if a.open and hasattr(os, "startfile"):
        os.startfile(path)


if __name__ == "__main__":
    main()
