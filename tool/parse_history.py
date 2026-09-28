"""
Leest de draftgeschiedenis uit het Google Sheets-export bestand en schrijft
data/draft_history.csv (year, team, name, price) plus data/alias_candidates.json.

    py tool\\parse_history.py "2025 Fantasy Auction Draft.xlsx"

Waarvoor: de 1456 werkelijk betaalde prijzen van deze league zijn de enige
marktdata die precies over deze twaalf managers gaat. Generieke ADP zegt wat de
markt doet, dit zegt wat jullie doen. Daarnaast levert het de aliastabel voor de
drafttool, want de namen worden tijdens de draft vrij getypt.

Layout van de tabbladen, over negen jaar in twee varianten:
  - teamnaam in de naamkolom, prijs in de kolom links ervan (2019 tot 2025)
  - teamnaam in de prijskolom, naam in de kolom rechts ervan (2017 en 2018)
De budgetrij is de eerste rij in de top 8 met minstens zes cellen op 200; de
teamnamen staan altijd een rij daarboven. De oriëntatie wordt per tabblad
bepaald door beide varianten te proberen en de variant met de meeste picks te
nemen, niet per kolom: per kolom beslissen telt kolommen dubbel.

Vereist openpyxl. Draaien in de map waar het bestand staat.
"""

import csv
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl ontbreekt. Installeer met: py -m pip install openpyxl")

YEAR = re.compile(r"(20\d\d)")
LABEL = re.compile(r"^(speler|budget|uitgegeven|resterend|max bid|totaal|owner|player)\b", re.I)
CLEANED_FROM = 2024  # vanaf dit jaar zijn de namen na de draft opgeschoond
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data")


def budget_row(ws):
    for r in range(1, 9):
        n = sum(1 for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value in (200, 200.0))
        if n >= 6:
            return r
    return None


def harvest(ws, brow, shift):
    hdr, cols, M = brow - 1, [], ws.max_column
    for c in range(1, M + 1):
        nm = ws.cell(hdr, c).value
        if not (nm and str(nm).strip()):
            continue
        pc, nc = (c - 1, c) if shift == 0 else (c, c + 1)
        if pc < 1 or nc > M or ws.cell(brow, pc).value not in (200, 200.0):
            continue
        cols.append((pc, nc, str(nm).strip()))

    picks = []
    for pc, nc, team in cols:
        for r in range(brow + 1, ws.max_row + 1):
            name, price = ws.cell(r, nc).value, ws.cell(r, pc).value
            if not name or not str(name).strip():
                continue
            s = str(name).strip()
            if LABEL.match(s) or not isinstance(price, (int, float)):
                continue
            picks.append({"team": team, "price": float(price), "name": s})
    return cols, picks


def parse_sheet(ws):
    brow = budget_row(ws)
    if not brow:
        return None
    best = None
    for shift in (0, 1):
        cols, picks = harvest(ws, brow, shift)
        if best is None or len(picks) > len(best[1]):
            best = (cols, picks)
    return best if best and best[1] else None


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "2025 Fantasy Auction Draft.xlsx"
    if not os.path.exists(path):
        raise SystemExit(f"Bestand niet gevonden: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)

    rows = []
    print(f"{'jaar':<7}{'teams':<7}{'picks':<7}{'gem/team':<10}duurste")
    for ws in wb.worksheets:
        t = ws.title
        if "prijzen" in t.lower():
            continue
        if not (t.startswith("Auction") or t.strip().isdigit()):
            continue
        m = YEAR.search(t)
        if not m:
            continue
        res = parse_sheet(ws)
        if not res:
            print(f"{m.group(1):<7}niet herkend, layout wijkt af")
            continue
        cols, picks = res
        top = max(picks, key=lambda q: q["price"])
        print(f"{m.group(1):<7}{len(cols):<7}{len(picks):<7}"
              f"{sum(q['price'] for q in picks) / len(cols):<10.0f}"
              f"{top['name']} ({top['price']:.0f})")
        rows += [{"year": int(m.group(1)), **q} for q in picks]

    if not rows:
        raise SystemExit("Geen picks gevonden.")

    os.makedirs(OUT, exist_ok=True)
    with open(os.path.join(OUT, "draft_history.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["year", "team", "name", "price"])
        w.writeheader()
        w.writerows(rows)

    # De aliastabel wordt op de ongekuiste jaren gebouwd. De opgeschoonde jaren
    # geven een te rooskleurig beeld van hoe netjes er tijdens de draft getypt wordt.
    raw = sorted({r["name"] for r in rows if r["year"] < CLEANED_FROM})
    alle = sorted({r["name"] for r in rows})
    json.dump({"ongekuist": raw, "alle": alle},
              open(os.path.join(OUT, "alias_candidates.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    odd = [n for n in raw
           if len(n.split()) < 2 or not re.match(r"^[A-Z][a-zA-Z'.\-]+ [A-Z]", n)]
    print(f"\n{len(rows)} picks over {len(set(r['year'] for r in rows))} jaar")
    print(f"aliaskandidaten uit de jaren voor {CLEANED_FROM}: {len(raw)} unieke strings, "
          f"{len(odd)} wijken af van 'Voornaam Achternaam' ({100 * len(odd) / max(len(raw), 1):.0f}%)")
    print(f"geschreven naar {os.path.normpath(OUT)}")


if __name__ == "__main__":
    main()
