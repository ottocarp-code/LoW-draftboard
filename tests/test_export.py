from datetime import datetime

import openpyxl
import pytest

import export_xlsx
from parse_history import parse_sheet

NOW = datetime(2026, 10, 20, 21, 15)


def pid(c, name):
    return next(p["id"] for p in c.get("/api/players").json()["players"] if p["name"] == name)


def draft_some(c, n=3):
    players = c.get("/api/players").json()["players"][:n]
    teams = ["RoRo", "RJ", "RoRo"]
    for p, t, price in zip(players, teams, (40, 12, 5)):
        r = c.post("/api/pick", json={"player_id": p["id"], "team": t, "price": price})
        assert r.status_code == 200, r.text
    return players


def test_export_reads_db_and_round_trips_through_history_parser(make_client, tmp_path):
    db = tmp_path / "draft.db"
    c = make_client(db=db)
    players = draft_some(c)

    path, draft = export_xlsx.export(str(db), str(tmp_path / "out"), now=NOW)
    assert len(draft["picks"]) == 3 and draft["teams"][0] == "RoRo"

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Auction 2026", "Picks"]
    cols, picks = parse_sheet(wb["Auction 2026"])
    assert len(cols) == 12
    got = sorted((q["team"], q["name"], q["price"]) for q in picks)
    want = sorted([("RoRo", players[0]["name"], 40.0), ("RJ", players[1]["name"], 12.0),
                   ("RoRo", players[2]["name"], 5.0)])
    assert got == want
    assert wb["Picks"].max_row == 4 and wb["Picks"]["C2"].value == players[0]["name"]


def test_export_formulas_for_spent_remaining_and_max_bid(make_client, tmp_path):
    db = tmp_path / "draft.db"
    draft_some(make_client(db=db))
    path, _ = export_xlsx.export(str(db), str(tmp_path / "out"), now=NOW)
    ws = openpyxl.load_workbook(path)["Auction 2026"]
    # 13 spots: picks in rows 4..16, footer from row 17, RoRo in columns A/B
    assert ws["B17"].value == "Uitgegeven" and ws["A17"].value == "=SUM(A4:A16)"
    assert ws["A18"].value == "=A3-A17"
    assert ws["A19"].value == "=IF((13-COUNT(A4:A16))<=0,0,MAX(0,A18-((13-COUNT(A4:A16))-1)))"


def test_export_works_while_the_app_holds_the_db_open(make_client, tmp_path):
    db = tmp_path / "draft.db"
    c = make_client(db=db)
    draft_some(c, 1)
    path, draft = export_xlsx.export(str(db), str(tmp_path / "out"), now=NOW)
    assert len(draft["picks"]) == 1
    assert len(c.get("/api/state").json()["picks"]) == 1   # the app still works afterwards


def test_template_copy_gets_current_teams_and_picks(make_client, tmp_path):
    db = tmp_path / "draft.db"
    c = make_client(db=db)
    players = draft_some(c)
    template, _ = export_xlsx.export(str(db), str(tmp_path / "tpl"), now=NOW)
    c.post("/api/undo")                                  # the template is now one pick ahead

    later = datetime(2027, 10, 1)
    path, _ = export_xlsx.export(str(db), str(tmp_path / "out"), template=template, now=later)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["Auction 2027", "Picks 2027"]
    assert "Auction 2026" in wb.sheetnames               # the template's own tab stays
    cols, picks = parse_sheet(wb["Auction 2027"])
    assert sorted((q["team"], q["name"]) for q in picks) == sorted(
        [("RoRo", players[0]["name"]), ("RJ", players[1]["name"])])
    assert wb["Auction 2027"]["A17"].value == "=SUM(A4:A16)"   # formulas kept


def test_missing_db_is_a_clear_error(tmp_path):
    with pytest.raises(SystemExit, match="niet gevonden"):
        export_xlsx.export(str(tmp_path / "nope.db"), str(tmp_path / "out"), now=NOW)
